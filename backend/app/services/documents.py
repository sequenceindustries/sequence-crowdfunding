"""Text extraction for uploaded pitch decks."""
import io


class ExtractionError(Exception):
    pass


def extract_text(data: bytes, ext: str):
    """Return (text, page_count). Raises ExtractionError on malformed files."""
    try:
        if ext == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise ExtractionError("Encrypted PDFs are not supported.")
            pages = [(p.extract_text() or "") for p in reader.pages]
            return "\n\n".join(pages).strip(), len(pages)
        if ext == ".pptx":
            from pptx import Presentation

            prs = Presentation(io.BytesIO(data))
            slides = []
            for slide in prs.slides:
                parts = []
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        parts.append(shape.text_frame.text)
                slides.append("\n".join(parts))
            return "\n\n".join(slides).strip(), len(slides)
    except ExtractionError:
        raise
    except Exception as exc:  # malformed or hostile input must never crash the worker
        raise ExtractionError(f"Could not read the document: {type(exc).__name__}") from exc
    raise ExtractionError("Unsupported file type.")
