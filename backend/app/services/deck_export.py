"""PPTX export of a deck draft."""
import io

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from .deck_builder import THEMES


def _rgb(hexstr):
    return RGBColor.from_string(hexstr)


def _textbox(slide, left, top, width, height):
    tf = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height)).text_frame
    tf.word_wrap = True
    return tf


def export_pptx(title: str, theme: str, slides: list) -> bytes:
    colors = THEMES[theme]
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    for s in slides:
        slide = prs.slides.add_slide(blank)
        bg = slide.background.fill
        bg.solid()
        bg.fore_color.rgb = _rgb(colors["bg"])
        is_cover = s["key"] == "cover"

        tf = _textbox(slide, 0.8, 2.4 if is_cover else 0.6, 11.7, 1.4)
        run = tf.paragraphs[0].add_run()
        run.text = s["title"]
        run.font.size = Pt(48 if is_cover else 34)
        run.font.bold = True
        run.font.color.rgb = _rgb(colors["title"])

        top = 3.9 if is_cover else 1.9
        tf = _textbox(slide, 0.8, top, 11.7, 4.6)
        first = True
        if s["body"]:
            p = tf.paragraphs[0]
            first = False
            r = p.add_run()
            r.text = s["body"]
            r.font.size = Pt(22)
            r.font.color.rgb = _rgb(colors["body"])
        for b in s["bullets"]:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.space_after = Pt(8)
            r = p.add_run()
            r.text = "•  " + b
            r.font.size = Pt(22 if len(s["bullets"]) <= 6 else 18)
            r.font.color.rgb = _rgb(colors["body"])

        if s.get("label"):
            tf = _textbox(slide, 0.8, 6.8, 11.7, 0.5)
            r = tf.paragraphs[0].add_run()
            r.text = s["label"]
            r.font.size = Pt(12)
            r.font.italic = True
            r.font.color.rgb = _rgb(colors["accent"])
        if s.get("incomplete"):
            tf = _textbox(slide, 0.8, 0.1, 11.7, 0.4)
            r = tf.paragraphs[0].add_run()
            r.text = "INCOMPLETE DRAFT: " + (s.get("note") or "this slide needs content")
            r.font.size = Pt(12)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0xB3, 0x26, 0x1E)

    prs.core_properties.title = title[:200]
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
