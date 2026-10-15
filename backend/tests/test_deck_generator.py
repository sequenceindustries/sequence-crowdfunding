import io
import re

from pptx import Presentation

FULL = {
    "business_name": "Sun Kiosks", "tagline": "Power for every commuter",
    "problem": "Commuters lose phone battery on long taxi trips.\nTaxi ranks have no safe charging.",
    "solution": "Solar charging kiosks at taxi ranks.", "product": "A kiosk with six charging ports.",
    "customers": "Daily taxi commuters.", "market": "Taxi ranks in Gauteng.",
    "competition": "Informal charging stalls.", "revenue_model": "Pay per charge.",
    "marketing": "Partner with taxi associations.", "operations": "Local technicians service the kiosks.",
    "traction": "One pilot kiosk running since March.",
    "financial_assumptions": "Assume 40 charges per day per kiosk.",
    "funding_request_cents": 1_000_000,
    "use_of_funds": [{"item": "Kiosk units", "amount_cents": 900_000}, {"item": "Installation", "amount_cents": 100_000}],
    "team": [{"name": "Thandi", "role": "Founder", "experience": "Ten years in energy"}],
    "milestones": [{"title": "Launch five kiosks", "date": "Q2"}],
    "risks": "Theft of equipment.",
}


def make(c, answers=None, **extra):
    r = c.post("/api/deck-drafts", {"answers": answers if answers is not None else FULL, **extra})
    assert r.status_code == 201, r.json
    return r.json["draft"]


def keys(d):
    return [s["key"] for s in d["slides"]]


def all_text(d):
    return " ".join(" ".join([s["title"], s["body"], *s["bullets"]]) for s in d["slides"])


def test_full_answers_produce_complete_ordered_deck(founder):
    d = make(founder)
    assert keys(d) == ["cover", "problem", "solution", "product", "market", "business_model", "competition",
                       "marketing", "traction", "team", "financials", "funding", "milestones", "risks"]
    assert d["incomplete_count"] == 0
    assert "Total: R10 000" in all_text(d)  # 1,000,000 cents


def test_slide_count_depends_on_the_project(founder):
    small = make(founder, {"business_name": "Tiny", "problem": "P", "solution": "S",
                           "funding_request_cents": 100, "use_of_funds": [{"item": "x", "amount_cents": 100}]})
    assert keys(small) == ["cover", "problem", "solution", "funding"]
    assert len(make(founder)["slides"]) > len(small["slides"])


def test_empty_answers_yield_only_flagged_essentials_no_invention(founder):
    d = make(founder, {})
    assert keys(d) == ["cover", "problem", "solution", "funding"]
    assert all(s["incomplete"] and s["note"] for s in d["slides"])
    assert d["incomplete_count"] == 4


def test_nothing_is_invented_for_missing_sections(founder):
    d = make(founder, {"business_name": "Acme", "problem": "Slow deliveries."})
    for absent in ("traction", "team", "financials", "market", "competition", "product", "milestones"):
        assert absent not in keys(d)
    assert d["incomplete_count"] == 2  # solution and funding are essential and missing


def test_no_number_appears_unless_the_founder_typed_it(founder):
    answers = {"business_name": "Acme", "problem": "Late deliveries cost small shops customers.",
               "solution": "A courier network.", "traction": "Pilot with a local shop."}
    d = make(founder, answers)
    given = set(re.findall(r"\d+", " ".join(str(v) for v in answers.values())))
    shown = set(re.findall(r"\d+", all_text(d)))
    assert shown <= given, f"invented figures: {shown - given}"


def test_every_word_on_content_slides_comes_from_the_founder(founder):
    d = make(founder)
    given = set(re.findall(r"[a-z']+", " ".join(
        str(v) if not isinstance(v, list) else " ".join(map(str, [x for i in v for x in i.values()])) for v in FULL.values()).lower()))
    scaffold = {"the", "problem", "our", "solution", "product", "or", "service", "customers", "and", "market",
                "business", "model", "competitive", "landscape", "marketing", "operations", "progress", "to", "date",
                "team", "financial", "assumptions", "funding", "request", "use", "of", "funds", "milestones",
                "risks", "next", "steps", "requested", "total", "r", "reported", "by", "owner", "not",
                "independently", "verified", "provided", "estimates", "planned", "owner-provided", "q"}
    for s in d["slides"]:
        words = set(re.findall(r"[a-z']+", " ".join([s["title"], s["body"], *s["bullets"], s["label"] or ""]).lower()))
        assert words <= given | scaffold, f"{s['key']}: {words - given - scaffold}"


def test_funding_total_mismatch_is_flagged_not_fixed(founder):
    a = {**FULL, "use_of_funds": [{"item": "Kiosk units", "amount_cents": 600_000}]}
    d = make(founder, a)
    f = [s for s in d["slides"] if s["key"] == "funding"][0]
    assert f["incomplete"] and "adds up to R6 000" in f["note"] and "requesting R10 000" in f["note"]
    assert "Total:" not in " ".join(f["bullets"])


def test_funding_missing_pieces_each_prompt_the_founder(founder):
    only_request = make(founder, {"funding_request_cents": 500_000})
    f = [s for s in only_request["slides"] if s["key"] == "funding"][0]
    assert f["incomplete"] and "use-of-funds" in f["note"]
    only_items = make(founder, {"use_of_funds": [{"item": "x", "amount_cents": 5}]})
    f = [s for s in only_items["slides"] if s["key"] == "funding"][0]
    assert f["incomplete"] and "total amount" in f["note"]


def test_assumptions_and_traction_labelled_unverified(founder):
    d = make(founder)
    by = {s["key"]: s for s in d["slides"]}
    assert "Not verified" in by["financials"]["label"] and "not independently verified" in by["traction"]["label"]


def test_edit_reorder_remove_and_add_slides(founder):
    d = make(founder)
    did = d["id"]
    slides = d["slides"]
    first, second = slides[0], slides[1]
    edited = [{**second, "title": "Why this matters"}, first] + slides[2:-1]  # swap, rename, drop the last
    edited.append({"title": "Thank you", "body": "Questions welcome.", "key": "custom"})
    r = founder.patch(f"/api/deck-drafts/{did}", {"slides": edited})
    assert r.status_code == 200
    out = r.json["draft"]["slides"]
    assert out[0]["title"] == "Why this matters" and out[0]["id"] == second["id"] and out[1]["id"] == first["id"]
    assert out[-1]["key"] == "custom" and out[-1]["id"] and "risks" not in [s["key"] for s in out]
    assert len(out) == len(slides)  # one removed, one added
    assert founder.get(f"/api/deck-drafts/{did}").json["draft"]["slides"] == out  # persisted


def test_clearing_a_slide_makes_it_incomplete(founder):
    d = make(founder)
    s = {**d["slides"][1], "bullets": [], "body": ""}
    r = founder.patch(f"/api/deck-drafts/{d['id']}", {"slides": [s]})
    assert r.json["draft"]["slides"][0]["incomplete"] is True
    s2 = {**d["slides"][1], "title": ""}
    assert founder.patch(f"/api/deck-drafts/{d['id']}", {"slides": [s2]}).json["draft"]["slides"][0]["incomplete"]


def test_unedited_mismatched_funding_slide_stays_flagged_after_round_trip(founder):
    d = make(founder, {**FULL, "use_of_funds": [{"item": "x", "amount_cents": 1}]})
    r = founder.patch(f"/api/deck-drafts/{d['id']}", {"slides": d["slides"]})  # client sends everything back unchanged
    f = [s for s in r.json["draft"]["slides"] if s["key"] == "funding"][0]
    assert f["incomplete"] is True and f["note"]


def test_slide_validation_rejects_bad_input(founder):
    did = make(founder)["id"]
    url = f"/api/deck-drafts/{did}"
    assert founder.patch(url, {"slides": "nope"}).status_code == 400
    assert founder.patch(url, {"slides": [1]}).status_code == 400
    assert founder.patch(url, {"slides": [{"id": "a", "title": "t", "bullets": "x"}]}).status_code == 400
    assert founder.patch(url, {"slides": [{"id": "a", "title": "t", "body": "b"}] * 2}).status_code == 400
    assert founder.patch(url, {"slides": [{"title": "t", "body": "b"}] * 41}).status_code == 400
    assert founder.patch(url, {"theme": "neon"}).status_code == 400
    assert founder.post("/api/deck-drafts", {"theme": "neon"}).status_code == 400
    assert founder.post("/api/deck-drafts", {"answers": {"problem": 5}}).status_code == 400
    assert founder.post("/api/deck-drafts", {"answers": {"funding_request_cents": -1}}).status_code == 400
    assert founder.post("/api/deck-drafts", {"answers": {"use_of_funds": [{"item": "x", "amount_cents": 1.5}]}}).status_code == 400


def test_regenerate_restores_section_from_answers(founder):
    d = make(founder)
    did = d["id"]
    prob = d["slides"][1]
    founder.patch(f"/api/deck-drafts/{did}", {"slides": [{**prob, "bullets": ["Totally rewritten"]}]})
    founder.patch(f"/api/deck-drafts/{did}", {"answers": {"problem": "A sharper problem statement."}})
    r = founder.post(f"/api/deck-drafts/{did}/regenerate", {"key": "problem"})
    assert r.status_code == 200
    s = r.json["draft"]["slides"][0]
    assert s["bullets"] == ["A sharper problem statement."] and s["id"] == prob["id"]


def test_regenerate_inserts_missing_section_in_order_and_rejects_empty(founder):
    d = make(founder, {"business_name": "A", "problem": "P", "solution": "S",
                       "funding_request_cents": 1, "use_of_funds": [{"item": "x", "amount_cents": 1}]})
    did = d["id"]
    assert founder.post(f"/api/deck-drafts/{did}/regenerate", {"key": "traction"}).status_code == 409
    founder.patch(f"/api/deck-drafts/{did}", {"answers": {"traction": "A pilot."}})
    r = founder.post(f"/api/deck-drafts/{did}/regenerate", {"key": "traction"})
    assert keys(r.json["draft"]) == ["cover", "problem", "solution", "traction", "funding"]
    assert founder.post(f"/api/deck-drafts/{did}/regenerate", {"key": "nonsense"}).status_code == 400


def test_answers_merge_without_touching_slides_until_regenerated(founder):
    d = make(founder)
    r = founder.patch(f"/api/deck-drafts/{d['id']}", {"answers": {"risks": "New risk."}})
    assert r.json["draft"]["answers"]["risks"] == "New risk." and r.json["draft"]["answers"]["business_name"] == "Sun Kiosks"
    assert [s for s in r.json["draft"]["slides"] if s["key"] == "risks"][0]["bullets"] == ["Theft of equipment."]


def deck_from(resp):
    return Presentation(io.BytesIO(resp.data))


def slide_text(prs):
    return [" ".join(sh.text_frame.text for sh in sl.shapes if sh.has_text_frame) for sl in prs.slides]


def test_export_complete_deck_is_valid_pptx_with_content(founder):
    d = make(founder, theme="midnight")
    r = founder.get(f"/api/deck-drafts/{d['id']}/export.pptx")
    assert r.status_code == 200 and r.headers["Cache-Control"] == "private, no-store"
    assert r.data[:2] == b"PK" and "pptx" in r.headers["Content-Disposition"]
    prs = deck_from(r)
    texts = slide_text(prs)
    assert len(texts) == 14 and "Sun Kiosks" in texts[0] and "Power for every commuter" in texts[0]
    assert any("Funding requested: R10 000" in t for t in texts)
    assert not any("INCOMPLETE" in t for t in texts)
    assert str(prs.slides[0].background.fill.fore_color.rgb) == "14181F"  # midnight theme applied


def test_incomplete_deck_cannot_be_exported_by_accident(founder):
    d = make(founder, {"business_name": "Acme", "problem": "Issue"})
    url = f"/api/deck-drafts/{d['id']}/export.pptx"
    r = founder.get(url)
    assert r.status_code == 409 and r.json["incomplete_slides"]
    forced = founder.get(url + "?allow_incomplete=true")
    assert forced.status_code == 200
    assert any("INCOMPLETE DRAFT" in t for t in slide_text(deck_from(forced)))


def test_empty_deck_blocked_and_control_characters_do_not_crash_export(founder):
    d = make(founder, {"business_name": "Bell\x07Co\x00", "problem": "Line\x0bbreak", "solution": "Fine",
                       "funding_request_cents": 5, "use_of_funds": [{"item": "x", "amount_cents": 5}]})
    assert "\x07" not in all_text(d)
    assert founder.get(f"/api/deck-drafts/{d['id']}/export.pptx").status_code == 200
    founder.patch(f"/api/deck-drafts/{d['id']}", {"slides": []})
    assert founder.get(f"/api/deck-drafts/{d['id']}/export.pptx").status_code == 409


def test_ownership_and_roles(founder, other_founder, client, make_client):
    did = make(founder)["id"]
    base = f"/api/deck-drafts/{did}"
    assert other_founder.get(base).status_code == 404
    assert other_founder.patch(base, {"title": "x"}).status_code == 404
    assert other_founder.post(base + "/regenerate", {"key": "problem"}).status_code == 404
    assert other_founder.get(base + "/export.pptx").status_code == 404
    assert other_founder.delete(base).status_code == 404
    assert other_founder.get("/api/deck-drafts").json["drafts"] == []
    assert client.get(base).status_code == 401
    sup = make_client()
    sup.register("s@example.com", role="supporter")
    assert sup.get(base).status_code == 403 and sup.post("/api/deck-drafts", {"answers": {}}).status_code == 403
    assert founder.delete(base).status_code == 200 and founder.get(base).status_code == 404


def test_draft_limit(app, founder):
    from app.api import deck_drafts

    deck_drafts.MAX_DRAFTS_PER_USER = 2
    try:
        make(founder, {}), make(founder, {})
        assert founder.post("/api/deck-drafts", {"answers": {}}).status_code == 429
    finally:
        deck_drafts.MAX_DRAFTS_PER_USER = 50
